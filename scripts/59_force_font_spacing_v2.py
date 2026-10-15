from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOCX = "manuscript.docx"
doc = Document(DOCX)

FONT = "Book Antiqua"
BODY_SIZE = Pt(12)
SPACING = 1.5

def is_heading_or_title(para):
    st = (para.style.name or "").lower()
    return st.startswith(("heading", "title", "subtitle"))

def force_run(run, size):
    run.font.name = FONT
    run.font.size = size
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), FONT)
    rFonts.set(qn('w:hAnsi'), FONT)
    rFonts.set(qn('w:cs'), FONT)
    rFonts.set(qn('w:eastAsia'), FONT)

def force_paragraph(para):
    if is_heading_or_title(para):
        # Headings: only set font family, keep heading size
        for run in para.runs:
            run.font.name = FONT
            rPr = run._element.get_or_add_rPr()
            rFonts = rPr.find(qn('w:rFonts'))
            if rFonts is None:
                rFonts = OxmlElement('w:rFonts')
                rPr.append(rFonts)
            rFonts.set(qn('w:ascii'), FONT)
            rFonts.set(qn('w:hAnsi'), FONT)
        para.paragraph_format.line_spacing = SPACING
    else:
        para.paragraph_format.line_spacing = SPACING
        for run in para.runs:
            force_run(run, BODY_SIZE)

def process_table(tbl):
    for row in tbl.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                force_paragraph(para)
            for nested in cell.tables:
                process_table(nested)

for para in doc.paragraphs:
    force_paragraph(para)

for tbl in doc.tables:
    process_table(tbl)

for section in doc.sections:
    for container in (section.header, section.footer,
                      section.first_page_header, section.first_page_footer,
                      section.even_page_header, section.even_page_footer):
        try:
            for para in container.paragraphs:
                force_paragraph(para)
        except Exception:
            pass

# Styles: only Normal/body styles get size 12; headings get font family only
for sname in ("Normal", "Body Text", "List Paragraph", "First Paragraph",
              "Table Grid", "Image Caption", "Caption", "No Spacing"):
    try:
        st = doc.styles[sname]
        st.font.name = FONT
        st.font.size = BODY_SIZE
        st.paragraph_format.line_spacing = SPACING
    except KeyError:
        pass

for sname in ("Heading 1", "Heading 2", "Heading 3", "Heading 4", "Title", "Subtitle"):
    try:
        st = doc.styles[sname]
        st.font.name = FONT
        st.paragraph_format.line_spacing = SPACING
    except KeyError:
        pass

doc.save(DOCX)
print("Applied Book Antiqua to body (12pt) and headings (font only). Headings retain their size hierarchy.")
