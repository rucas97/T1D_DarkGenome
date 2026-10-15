# scripts/61_style_all_tables.py
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOCX = "manuscript.docx"
doc = Document(DOCX)

def clear_and_set(cell, text, size=9, bold=False, align='left'):
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)
    p = cell.add_paragraph()
    r = p.add_run(str(text))
    r.font.name = "Book Antiqua"
    r.font.size = Pt(size)
    r.bold = bold
    p.paragraph_format.line_spacing = 1.0
    if align == 'center':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'right':
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

def set_col_width(cell, width_in):
    cell.width = Inches(width_in)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn('w:tcW'))
    if tcW is None:
        tcW = OxmlElement('w:tcW')
        tcPr.append(tcW)
    tcW.set(qn('w:w'), str(int(width_in * 1440)))
    tcW.set(qn('w:type'), 'dxa')

def set_row_height(row, inches, rule='atLeast'):
    trPr = row._tr.get_or_add_trPr()
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), str(int(inches * 1440)))
    trHeight.set(qn('w:hRule'), rule)
    trPr.append(trHeight)

def style_table(tbl, col_widths, header_aligns, body_aligns,
                header_size=10, body_size=9, header_row_height=0.5):
    n_cols = len(col_widths)

    # Header row
    if tbl.rows:
        hdr = tbl.rows[0]
        for i, cell in enumerate(hdr.cells):
            if i >= n_cols:
                break
            # Preserve existing text, restyle
            txt = cell.text.strip()
            align = header_aligns[i] if i < len(header_aligns) else 'center'
            clear_and_set(cell, txt, size=header_size, bold=True, align=align)
            set_col_width(cell, col_widths[i])
        set_row_height(hdr, header_row_height, rule='atLeast')

    # Body rows
    for row in tbl.rows[1:]:
        for i, cell in enumerate(row.cells):
            if i >= n_cols:
                break
            txt = cell.text.strip()
            align = body_aligns[i] if i < len(body_aligns) else 'left'
            clear_and_set(cell, txt, size=body_size, bold=False, align=align)
            set_col_width(cell, col_widths[i])

# ============================================================
# Table 1 — Dataset | Design | Role (3 cols)
# ============================================================
if len(doc.tables) >= 1:
    style_table(
        doc.tables[0],
        col_widths=[1.35, 3.15, 2.00],
        header_aligns=['left', 'left', 'left'],
        body_aligns=['left', 'left', 'left'],
    )
    print("Table 1 styled.")

# ============================================================
# Table 2 — Target | log2FC | Neighbor (3 cols)
# ============================================================
if len(doc.tables) >= 2:
    style_table(
        doc.tables[1],
        col_widths=[1.80, 0.90, 3.80],
        header_aligns=['left', 'center', 'left'],
        body_aligns=['left', 'center', 'left'],
    )
    print("Table 2 styled.")

# ============================================================
# Table 3 — Donor ID | control | treated | log2FC (4 cols)
# ============================================================
if len(doc.tables) >= 3:
    style_table(
        doc.tables[2],
        col_widths=[1.10, 1.75, 1.75, 1.90],
        header_aligns=['center', 'center', 'center', 'center'],
        body_aligns=['center', 'right', 'right', 'center'],
    )
    print("Table 3 styled.")

# ============================================================
# Table 4 — already styled, re-apply for consistency (8 cols)
# ============================================================
if len(doc.tables) >= 4:
    style_table(
        doc.tables[3],
        col_widths=[1.45, 0.85, 0.55, 0.85, 0.65, 0.65, 0.50, 0.50],
        header_aligns=['center'] * 8,
        body_aligns=['left'] + ['center'] * 7,
    )
    print("Table 4 re-styled.")

# ============================================================
# Table 5 — Peptide | Allele | MixMHC2pred | NetMHCIIpan | IC50_nM | Binding_core (6 cols)
# ============================================================
if len(doc.tables) >= 5:
    style_table(
        doc.tables[4],
        col_widths=[1.50, 0.95, 0.85, 0.85, 0.75, 1.10],
        header_aligns=['center'] * 6,
        body_aligns=['left', 'center', 'center', 'center', 'center', 'center'],
    )
    print("Table 5 styled.")

doc.save(DOCX)
print("Saved:", DOCX)